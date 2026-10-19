"""Validate xlsx files for common agent-generated delivery defects."""

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer


EMPTY_CACHED_RE = re.compile(rb"</f>\s*<v\s*/>|</f>\s*<v>\s*</v>")
CELL = r"(?:'[^']+'|[A-Za-z_][A-Za-z0-9_ .]*)?!?\$?[A-Z]{1,3}\$?\d+"
CELL_RANGE_RE = re.compile(rf"^{CELL}(?::{CELL})?$", re.IGNORECASE)
COL_RANGE_RE = re.compile(
    r"^(?:'[^']+'|[A-Za-z_][A-Za-z0-9_ .]*)?!?\$?[A-Z]{1,3}:\$?[A-Z]{1,3}$",
    re.IGNORECASE,
)
ROW_RANGE_RE = re.compile(
    r"^(?:'[^']+'|[A-Za-z_][A-Za-z0-9_ .]*)?!?\$?\d+:\$?\d+$",
    re.IGNORECASE,
)
ERROR_LITERALS = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}
LOGICAL_LITERALS = {"TRUE", "FALSE"}
PROSE_SYMBOLS = {
    "\u00d7": "multiplication sign",
    "\u00f7": "division sign",
    "\u00a7": "section sign",
    "\u2192": "arrow",
    "\u2013": "en dash",
    "\u2014": "em dash",
    "\u2248": "approximation sign",
}


def strip_quoted_text(value):
    return re.sub(r'"(?:[^"]|"")*"', '""', value)


def is_known_range(value, defined_names, table_names):
    token = value.strip()
    if not token:
        return True
    upper = token.upper()
    if upper in ERROR_LITERALS or upper in LOGICAL_LITERALS:
        return True
    if token in defined_names or token.split("[", 1)[0] in table_names:
        return True
    if "[" in token and "]" in token:
        return True
    return bool(
        CELL_RANGE_RE.match(token)
        or COL_RANGE_RE.match(token)
        or ROW_RANGE_RE.match(token)
    )


def suspicious_formula_reasons(formula, defined_names, table_names):
    body = formula[1:]
    reasons = []
    if not body.strip():
        reasons.append("empty formula")
        return reasons
    if body[0].isspace():
        reasons.append("whitespace immediately after '='; prose notes must not start with '='")

    unquoted = strip_quoted_text(body)
    symbols = [name for char, name in PROSE_SYMBOLS.items() if char in unquoted]
    if symbols:
        reasons.append("unquoted prose symbol(s): " + ", ".join(symbols))

    bad_operands = []
    try:
        tokens = Tokenizer(formula).items
    except Exception as exc:
        reasons.append(f"formula tokenizer failed: {exc}")
        return reasons

    for token in tokens:
        if token.type == "OPERAND" and token.subtype == "RANGE":
            value = token.value.strip()
            if not is_known_range(value, defined_names, table_names) and re.search(
                r"[A-Za-z]", value
            ):
                bad_operands.append(value)

    if len(bad_operands) >= 2:
        sample = ", ".join(dict.fromkeys(bad_operands[:8]))
        reasons.append("unrecognized word-like formula operands: " + sample)
    return reasons


def validate(path, allow_em_dash=False):
    path = Path(path)
    report = {
        "path": str(path),
        "status": "success",
        "formula_count": 0,
        "errors": [],
        "warnings": [],
    }

    if not path.exists():
        report["errors"].append(f"file does not exist: {path}")
        report["status"] = "failed"
        return report

    try:
        with zipfile.ZipFile(path) as archive:
            bad = archive.testzip()
            if bad:
                report["errors"].append(f"zip CRC check failed at {bad}")

            for item in archive.infolist():
                name = item.filename.replace("\\", "/")
                data = archive.read(item.filename)
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    if EMPTY_CACHED_RE.search(data):
                        report["errors"].append(
                            f"{name}: empty cached formula value remains; run finalize_xlsx.py"
                        )
                if name.endswith((".xml", ".rels")):
                    try:
                        ElementTree.fromstring(data)
                    except Exception as exc:
                        report["errors"].append(f"{name}: malformed XML: {exc}")
                    if not allow_em_dash:
                        text = data.decode("utf-8", errors="ignore")
                        if "\u2014" in text or "&#8212;" in text:
                            report["errors"].append(f"{name}: contains em dash U+2014")
    except Exception as exc:
        report["errors"].append(f"cannot read xlsx package: {exc}")
        report["status"] = "failed"
        return report

    try:
        wb = load_workbook(path, data_only=False)
    except Exception as exc:
        report["errors"].append(f"openpyxl cannot load workbook: {exc}")
        report["status"] = "failed"
        return report

    defined_names = set(wb.defined_names.keys())
    table_names = set()
    for ws in wb.worksheets:
        table_names.update(ws.tables.keys())

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    report["formula_count"] += 1
                    reasons = suspicious_formula_reasons(value, defined_names, table_names)
                    if reasons:
                        report["errors"].append(
                            f"{ws.title}!{cell.coordinate}: suspicious formula {value!r} ({'; '.join(reasons)})"
                        )
                elif isinstance(value, str) and value.startswith("'="):
                    report["warnings"].append(
                        f"{ws.title}!{cell.coordinate}: literal text begins with '=' via apostrophe"
                    )

    if report["errors"]:
        report["status"] = "failed"
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("path")
    parser.add_argument("--allow-em-dash", action="store_true")
    args = parser.parse_args()
    report = validate(args.path, allow_em_dash=args.allow_em_dash)
    print(json.dumps(report, indent=2))
    return 1 if report["status"] != "success" else 0


if __name__ == "__main__":
    sys.exit(main())
